from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent / "电子技术课程设计总结报告_视觉自主导航与无线体感控制_正式版.docx"

def font(run, name="宋体", size=None, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill); pr.append(node)

def body(text):
    p = doc.add_paragraph(text)
    for r in p.runs: font(r)
    return p

def heading(text, level=1):
    return doc.add_heading(text, level=level)

def table(headers, rows, sizes=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True
    hdr = t.rows[0]
    repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true")
    hdr._tr.get_or_add_trPr().append(repeat)
    for i, x in enumerate(headers):
        hdr.cells[i].text = str(x); shade(hdr.cells[i], "D9EAF7")
        hdr.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hdr.cells[i].paragraphs[0].runs: font(r, "黑体", 9, True)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = str(x); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n % 2: shade(cells[i], "F7FAFC")
            if sizes: cells[i].width = Cm(sizes[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1.15
                for r in p.runs: font(r, "宋体", 9)
    doc.add_paragraph()
    return t

def code(title, text):
    p = doc.add_paragraph(title); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    for r in p.runs: font(r, "黑体", 10, True)
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; shade(t.cell(0, 0), "F3F4F6")
    p = t.cell(0, 0).paragraphs[0]; p.style = doc.styles["Code"]
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text.strip()); font(r, "Consolas", 8.2)

def formula(no, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text + f"    （{no}）"); font(r, "Cambria Math", 11)

def field(paragraph, instruction, fallback=""):
    run = paragraph.add_run(); begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = fallback
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, text, end])

doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2.4), Cm(2.2), Cm(2.7), Cm(2.3)
normal = doc.styles["Normal"]; normal.font.name = "宋体"; normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.paragraph_format.line_spacing = 1.5; normal.paragraph_format.first_line_indent = Cm(0.74)
normal.paragraph_format.space_after = Pt(0)
for name, fn, sz in [("Title","黑体",22),("Heading 1","黑体",16),("Heading 2","黑体",14),("Heading 3","黑体",12)]:
    s=doc.styles[name]; s.font.name=fn; s.font.size=Pt(sz); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
    s._element.rPr.rFonts.set(qn("w:eastAsia"),fn); s.paragraph_format.keep_with_next=True
if "Code" not in [s.name for s in doc.styles]:
    s=doc.styles.add_style("Code",WD_STYLE_TYPE.PARAGRAPH); s.font.name="Consolas"; s.font.size=Pt(8.2)
    s._element.rPr.rFonts.set(qn("w:eastAsia"),"等线"); s.paragraph_format.line_spacing=1
hp=sec.header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=hp.add_run("电子技术课程设计总结报告"); font(r,"宋体",9)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; field(fp," PAGE ")

# 封面、摘要与目录
for _ in range(3): doc.add_paragraph()
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("电子技术课程设计总结报告"); font(r,"黑体",26,True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(28)
r=p.add_run("基于 ESP32-S3 的视觉自主导航与\n无线体感控制智能小车设计"); font(r,"黑体",20,True)
doc.add_paragraph()
table(["项目","内容"],[("课程名称","电子技术课程设计"),("院系专业","待填写"),("班级","待填写"),("学生姓名","王润宇、何晓瑜、王铜锌"),("学号","待填写"),("指导教师","待填写"),("完成日期","待填写")],[4,9])
doc.add_page_break(); heading("摘  要")
body("本课题设计并实现了一套以ESP32-S3为主控制器、兼具自主导航和无线体感控制能力的三轮全向智能小车。项目沿“物理红外循迹与超声波避障—机器视觉自主循迹与推球—STM32无线体感控制”三条技术路线递进。基础阶段采用四位掩码、PD方向修正、动态降速、三轮运动学混控与编码器PI闭环实现循迹，并以连续测距确认、数据时效检查和编码器状态机完成横移避障。视觉阶段将摄像头同一水平线上的四块采样区抽象为伪红外通道，复用已有循迹控制器；通过自适应阈值、终点三帧确认、正常转弯与丢线搜索双机制提升复杂线路适应性。推球任务对红球、蓝球和洞口进行连通域与几何筛选，十四状态控制器中的FAR_SPRINT远场冲刺策略把视觉对准、编码器定距、近场横移和最终直推结合起来。创新体感系统采用STM32F103C8T6读取MPU6050加速度静态倾角，经低通、滞回和连续帧防抖生成动作命令，再经ZS-040 BLE链路发送给车端；车端设置200 ms通信看门狗并由LCD反馈状态。系统形成了感知、决策、运动、通信与显示闭环，具有清晰的软件边界和多层安全失效机制。")
p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Cm(0)
r=p.add_run("关键词："); font(r,"黑体",11,True)
r=p.add_run("ESP32-S3；机器视觉；伪红外；循迹避障；推球状态机；STM32；无线体感控制"); font(r)
doc.add_page_break(); heading("目  录")
field(doc.add_paragraph(),' TOC \\o "1-3" \\h \\z \\u ',"打开文档后右键选择“更新域”生成目录页码。")
doc.add_page_break()

# 1 课题内容与总体方案
heading("1 课题内容与总体技术方案")
heading("1.1 课题内容",2)
body("本课题要求智能车沿黑色赛道自主行驶，识别并绕过障碍物，利用摄像头寻找指定颜色小球并推入对应洞口；本组还在课程任务之上开发无线体感控制。老师统一提供底盘、电机驱动和常规传感硬件，因此本文不把公共硬件作为设计成果展开，只在算法需要处说明接口，重点讨论数字系统的软件架构、感知算法、运动控制、通信协议和调试过程。")
body("系统有自主和体感两种模式。自主模式由ESP32-S3完成红外/超声波/摄像头采样、决策和三轮运动控制；体感模式由STM32读取手持端姿态，生成F、B、L、R、S五种单字符命令，经BLE发送到车端ESP32-S3。两种模式最终都使用明确的运动输出，并把关键状态显示在LCD、串口或Wi-Fi网页上。")
heading("1.2 技术演进路线",2)
table(["阶段","核心任务","主要方法","成果"],[
    ("一","物理红外循迹与超声波避障","四位掩码、PD、编码器PI、避障状态机","建立稳定运动基线"),
    ("二","视觉循迹、找球与推球","伪红外、连通域筛选、十四状态控制器","完成课程自主任务"),
    ("三","STM32无线体感控制","静态倾角、滤波/滞回/防抖、BLE看门狗","完成自主创新功能")])
heading("1.3 总体数据流与设计原则",2)
body("自主模式数据流为：感知输入→有效性与时效检查→离散掩码或目标几何量→循迹、避障或推球状态机→前进、横移、转向三个抽象运动量→三轮混控→轮级PI→PWM。体感模式数据流为：MPU6050加速度→俯仰/横滚角→低通滤波→手势判决→UART/BLE→车端命令解析→通信看门狗→电机输出。控制层以10 ms为主节拍运行，慢速感知和显示利用时间戳与快照接入。")
body("设计遵循三项原则：感知先转为可复用中间表示；长动作必须拆分为状态并设置超时、堵转和数据过期保护；阈值、状态跳转与电机响应必须能从串口、LCD或Wi-Fi页面观察，从而提高调试可解释性。")

# 2 模块与分工
heading("2 系统模块划分及队员工作分工")
heading("2.1 模块划分",2)
table(["模块","主要内容","输出"],[
    ("感知","红外、超声波、摄像头、MPU6050的采样、滤波、目标提取和时效判断","掩码、距离、目标位置、姿态角"),
    ("决策","循迹、转弯、丢线搜索、避障、推球和手势判决","forward/lateral/turn或字符命令"),
    ("运动","三轮运动学映射、限幅、编码器PI和堵转检查","三路PWM与方向"),
    ("通信","Wi-Fi视频/状态、BLE扫描订阅、UART与心跳","网页状态和车端命令"),
    ("显示","电机、距离、视觉、连接和动作状态格式化","LCD、串口和手机页面")])
heading("2.2 队员工作分工",2)
table(["成员","主要分工","具体工作"],[
    ("王润宇","循迹与避障代码开发","物理红外/伪红外接口、PD、弯道与丢线、超声波避障状态机及调试"),
    ("何晓瑜","推球代码开发、硬件调试","红蓝球与洞口识别、推球控制、远场冲刺策略和硬件排查"),
    ("王铜锌","STM32部分开发、硬件调试","MPU6050采集、手势判决、UART/BLE、接收端与LCD反馈"),
    ("全体成员","共同工作","整车安装、环境配置、联调、参数标定、验收、视频、汇报与报告")])
body("分工表示主要责任而非相互割裂。三个子系统最终共享电机、显示和通信资源，因此关键接口、实车参数标定和验收由全体成员共同完成。")

# 3 自主循迹与避障
heading("3 自主循迹与避障系统设计")
heading("3.1 四路红外与统一掩码",2)
body("四路红外编码为低四位掩码：bit0为车辆右外侧CH1，bit1为右内侧CH2，bit2为左内侧CH3，bit3为左外侧CH4。连续循迹误差只使用内侧两路，外侧两路仅用于转弯候选，避免宽线和缓弯造成过早强转。统一掩码也使第四章的摄像头伪红外能够复用本控制器。")
code("代码3-1 内侧通道误差映射（line_follow.c）",'''static int mask_to_error(uint8_t mask)
{
    const bool right = (mask & IR_CHANNEL_2_MASK) != 0U;
    const bool left  = (mask & IR_CHANNEL_3_MASK) != 0U;
    if (right == left) return 0;
    return right ? 1 : -1;
}''')
heading("3.2 PD循迹与动态降速",2)
body("比例项根据当前偏差纠正方向，微分项根据误差变化抑制摆动。基础前进量为170，检测到偏差时降低20且不低于120，转向量限幅±250。该策略在直线保持速度，在弯道主动留出转向余量。")
formula("3-1","u(k)=70e(k)+15[e(k)−e(k−1)]")
code("代码3-2 PD修正核心",'''int derivative = error - controller->previous_error;
int turn = clamp_int(LINE_KP * error + LINE_KD * derivative,
                     -LINE_TURN_LIMIT, LINE_TURN_LIMIT);
int forward = LINE_BASE_FORWARD;
if (error != 0) forward = max(LINE_MIN_FORWARD,
                              forward - LINE_CURVE_SLOWDOWN);''')
heading("3.3 三轮运动学与轮级PI",2)
body("上层只给出前进F、横移S和转向T，不直接写三路PWM。三轮命令按式（3-2）计算；若任一轮绝对值超过1000，则按峰值等比例缩放全部命令，保持运动方向。循迹主要使用F+T，避障使用F+S，推球可组合三个自由度。")
formula("3-2","L=F+T−S/2；R=F−T+S/2；B=−S")
body("轮级PI每50 ms更新。目标速度按每秒计数=2×命令换算，测速经(3×旧值+新值)/4平滑，Kp=1/2、Ki=1/4，积分限幅±1200、输出限幅500，并设置启动/运行最小驱动力和反积分饱和，解决低占空比不转及负载不一致。")
formula("3-3","PWM(k)=sat{Kp[n*(k)−n(k)]+KiΣ[n*−n]}")
heading("3.4 超声波连续确认与横移避障",2)
body("超声波依据d=cΔt/2计算距离。小于等于50 mm连续2次才触发绕障，大于120 mm连续3次才确认通道恢复；150 mm以内先把前进量降至90。测距超过500 ms未更新即判为无效，防止使用陈旧数据。")
formula("3-4","d=c·Δt/2")
body("避障由刹车、左横移、前进越障、右横移回线和结束状态组成。触发后制动100 ms，以120横移；左移至少240计数，前进目标1000计数、速度140，右移额外增加600计数并结合红外居中提前结束，清障条件需持续250 ms。编码器定距比固定延时更能适应电量和地面差异。")
heading("3.5 防锁死机制",2)
body("每个运动阶段设置6000 ms总超时；若500 ms内三轮总进度不足2计数，则判为堵转；累计进度超过合理上限也判为异常。任何保护触发都立即停车，覆盖车轮卡住、传感器过期和状态无法退出三类风险。")
code("代码3-3 避障动作失效保护（obstacle_avoidance.c）",'''if (progress - last_motion >= AVOID_STALL_MIN_COUNTS) {
    last_motion = progress; stall_ms = 0;
} else stall_ms += elapsed_ms;
return stage_ms >= AVOID_MOTION_TIMEOUT_MS ||
       stall_ms >= AVOID_STALL_TIMEOUT_MS ||
       (maximum_counts > 0 && progress > maximum_counts);''')

# 4 视觉导航
heading("4 基于机器视觉的自主导航系统设计")
heading("4.1 四块伪红外采样",2)
body("视觉循迹是课程要求，本组的创新在于把二维图像抽象为四路伪红外。LINE模式把480×320 MJPEG按1/8尺度解码为60×40，在约70%图像高度处放置四个4×4采样框。中心从车辆右侧到左侧依次位于宽度65%、54%、47%和35%，输出位序与物理红外CH1—CH4一致，因此无需重写PD、弯道和回线逻辑。")
table(["通道","图像位置","掩码位","用途"],[
    ("CH1","宽度65%，车辆右外侧","bit0","右弯候选"),("CH2","宽度54%，右内侧","bit1","误差+1"),
    ("CH3","宽度47%，左内侧","bit2","误差−1"),("CH4","宽度35%，车辆左外侧","bit3","左弯候选")])
heading("4.2 自适应阈值与终点确认",2)
body("系统先估计图像背景亮度，再以背景亮度减40作为黑线阈值；单框至少出现2个暗像素才置位。相对阈值比固定灰度更能适应室内灯光和曝光变化。全黑可能来自终点、阴影或瞬时模糊，故连续3帧全黑才锁存终点。四框采样值、阈值和掩码均输出到Wi-Fi页面，便于区分感知错误与控制参数错误。")
heading("4.3 正常转弯与丢线搜索",2)
body("正常转弯利用外侧通道预告直角弯：外侧命中持续30 ms后进入候选，在450 ms窗口内等待内侧和线形确认；确认后以200原地转动，最少60 ms，重新居中持续30 ms后退出，单次最长600 ms并保持50 ms退出过程。若四路全白，则进入编码器定距搜索：先向右约640计数（约120°），未找到后向左约1280计数（约240°），粗转230、精调180，回线稳定30 ms后交还PD；同时设置400 ms堵转和5 s总超时。两套机制分别处理“仍看见弯道”和“已经丢线”。")
heading("4.4 视觉流水线实时性",2)
body("USB回调只复制MJPEG并唤醒任务，软件JPEG解码放在core 1专用任务中，避免阻塞控制环和Wi-Fi。LINE使用1/8尺度，PUSH为获得几何精度使用1/4尺度。解码忙时丢弃新帧，不让旧帧排队；循迹帧超过1500 ms、找球帧超过500 ms便停止使用。该优化表述为降低延迟与提高有效处理速度，不宣称最终固定50 fps；最终输入配置为480×320、25 fps。")
code("代码4-1 解码任务配置（camera_vision.c）",'''#define VISION_TASK_STACK_SIZE 8192
#define VISION_TASK_PRIORITY   1
#define VISION_TASK_CORE       1
/* callback only copies MJPEG and wakes decoder;
   a new frame is dropped while decoding. */''')

# 5 推球
heading("5 视觉找球与自动推球系统设计")
heading("5.1 红球、蓝球与洞口检测",2)
body("PUSH模式按1/4尺度解码为120×80。红球候选满足R−G≥25且R−B≥25；蓝球候选满足B−R≥25、B−G≥25、B≥50且饱和度≥45。二值图经连通域分析，再按面积、紧致度、宽高比、长短轴、中心位置去除赛道直角、阴影和细长反光。蓝色候选还限制中心纵坐标不超过66并要求亮度梯度不低于3。洞口按远处低亮、低饱和连通域提取，并结合面积、边界和左右分区得到左右洞口。任务顺序为红球—左洞口、蓝球—右洞口。")
heading("5.2 十四状态控制器",2)
table(["状态","作用","关键条件"],[
    ("START_FORWARD","离开起始区","速度150、650 ms"),("FIND_BALL","寻找当前目标","看见球则接近"),
    ("APPROACH_BALL","横移居中、远场修正和接近","中心误差、半径、纵坐标"),("SCAN","左右定距扫描","找到目标或失败"),
    ("ALIGN","只横移使球与洞口对齐","误差满足300 ms"),("BACKOFF","留出冲刺空间","130、220 ms"),
    ("FAR_SPRINT","编码器定距冲刺","目标200～1200计数"),("PUSH","最终直线强推","300、650 ms"),
    ("BACKOUT","退出近场","计时/进度"),("EGRESS","离开目标区","120、1000 ms"),
    ("POST_EGRESS_TURN","约120°转向下一目标","编码器"),("POST_EGRESS_FORWARD","前进到下一搜索区","1000计数、速度80"),
    ("DONE","两球完成停车","终态"),("FAULT_STOP","异常停车","超时、堵转或数据失效")])
body("总任务超时180 s。每个编码器动作记录开始计数、当前进度和最后运动时刻，命令存在而计数不增长时进入FAULT_STOP；终态在同一周期立即发布，避免主循环锁止后网页仍显示旧状态。")
heading("5.3 FAR_SPRINT远场冲刺射门",2)
body("“远程射门”在本项目中不是网络遥控，而是FAR_SPRINT远场冲刺。球半径≤5、纵坐标≤35且球—洞口横向误差≤6时，系统在远场先完成对准，再将球纵向位置映射为200～1200计数，以200直线冲刺。近场ALIGN只用35～45横移，不前进、不旋转；误差≤8并持续300 ms后以300强推650 ms。该分段方法利用远场视觉和近场里程信息，缓解球进入车头盲区后闭环失效。")
code("代码5-1 FAR_SPRINT触发逻辑摘要",'''if (far_sprint_enable && ball->radius <= 5 && ball->cy <= 35 &&
    abs(pocket_error) <= 6) {
    counts = clamp(distance_from_y(ball->cy), 200, 1200);
    reset_stage(c, BALL_PUSH_FAR_SPRINT, counts, left, right);
}''')

# 6 STM32体感
heading("6 基于STM32与IMU的无线体感控制系统设计")
heading("6.1 接口与采样配置",2)
body("手持端采用STM32F103C8T6，I²C1的PB6/PB7分别连接SCL/SDA，以100 kHz访问地址0x68的MPU6050。程序扫描0x01～0x77并读取WHO_AM_I确认器件。配置为125 Hz、DLPF加速度44 Hz/陀螺仪42 Hz、±2g和±250°/s。一次读取14字节虽包含陀螺仪数据，但当前姿态只使用加速度静态倾角，不能写成六轴融合。")
heading("6.2 倾角、滤波与手势判决",2)
formula("6-1","pitch=atan2(−ax, √(ay²+az²))·180/π")
formula("6-2","roll=atan2(ay, az)·180/π")
formula("6-3","θf(k)=0.3θ(k)+0.7θf(k−1)")
body("一阶低通系数α=0.3。动作进入阈值30°、退出阈值20°，形成滞回；俯仰和横滚同时超限时选择绝对值更大的优势轴，映射为F/B/L/R，回中为S。候选必须连续3帧一致才生效。循环每100 ms发送一次，因此确认约300 ms。")
code("代码6-1 姿态计算与低通（STM32 main.c）",'''pitch = atan2f(-ax, sqrtf(ay*ay + az*az)) * 57.29578f;
roll  = atan2f( ay, az) * 57.29578f;
pitch_f = 0.3f*pitch + 0.7f*pitch_f;
roll_f  = 0.3f*roll  + 0.7f*roll_f;''')
heading("6.3 UART、BLE与失联停车",2)
body("USART1使用PA9/PA10、9600 baud，每100 ms发送单字节F/B/L/R/S。115200 baud曾导致ZS-040乱码，统一到9600后恢复。ESP32-S3作为BLE GATT客户端扫描并按MAC匹配，连接0xFFE0服务，枚举可Notify特征并写CCCD订阅。Notify只接受五种合法字符并更新时间戳；断开或200 ms无新命令立即停车。前进速度170、转向110，当前体感版本后轮保持0。")
code("代码6-2 Notify解析与通信看门狗（逻辑摘要）",'''if (length && strchr("FBLRS", data[0])) {
    command = data[0]; last_rx_ms = millis();
}
if (!connected || millis() - last_rx_ms > 200) {
    drive(0, 0, 0);
    showState(connected ? "LOST" : "NO LINK");
} else executeCommand(command);''')
heading("6.4 LCD反馈与可靠性",2)
body("ST7735每200 ms显示左右轮、后轮输出和NO LINK、LOST、FWD、BACK、LEFT、RIGHT、STOP状态，串口每1 s输出心跳。低通、滞回、优势轴、三帧防抖解决输入抖动；周期发送、200 ms看门狗解决链路异常；LCD使连接和动作可观察，构成完整安全链。")

# 7 软件架构
heading("7 数字系统软件架构及关键代码实现")
heading("7.1 三类开发框架与四套工程",2)
table(["工程","平台","职责","关键文件"],[
    ("lnf_avoid","ESP-IDF/ESP32-S3","物理红外循迹、超声波避障和运动基线","line_follow、obstacle_avoidance、drive_control"),
    ("vison_lnf_avoid","ESP-IDF/ESP32-S3","伪红外、找球推球、Wi-Fi画面与状态","camera_vision、ball_vision、ball_push"),
    ("STM32_LQ6050","STM32Cube/STM32F103C8T6","IMU、手势判决、UART发送","src/main.c"),
    ("ESP32_Car_Arduino","Arduino/ESP32-S3","BLE客户端、电机和LCD","src/main.cpp")])
body("前两个工程是自主任务的递进版本；后两个工程分别是体感手持端和车端接收验证程序。模块通过掩码、距离快照、视觉结果、抽象运动量和字符命令交换数据，不让感知算法直接依赖PWM硬件细节。")
heading("7.2 多速率主循环",2)
body("自主控制以10 ms为基本周期：读取模式、感知快照和编码器，按安全停车、避障、推球、循迹的优先级决策，最后一次性提交电机命令。轮级PI以50 ms更新，显示和网页以更慢周期刷新；视觉任务异步写带时间戳快照，控制循环不等待摄像头。")
code("代码7-1 跨模块控制主线（结构化伪代码）",'''every 10 ms:
    sensor, encoder = acquire_snapshots()
    if not sensor.fresh:       cmd = STOP
    elif avoidance.active:     cmd = obstacle_update(sensor, encoder)
    elif mode == PUSH:         cmd = ball_push_update(sensor, encoder)
    else:                      cmd = line_follow_update(sensor.mask)
    wheel = drive_mix_motion(cmd.forward, cmd.lateral, cmd.turn)
    motor_apply(wheel)''')
heading("7.3 文件职责与验证",2)
table(["模块","承担的职责","刻意不承担的职责"],[
    ("pseudo_infrared","图像采样、阈值和掩码","不决定电机动作"),("line_follow","掩码到循迹/转弯/搜索命令","不访问摄像头"),
    ("obstacle_avoidance","绕障状态与安全保护","不直接写PWM"),("ball_vision","红蓝球和洞口候选","不决定推球阶段"),
    ("ball_push","十四状态任务决策","不实现图像分割"),("drive_control/motor","运动学、闭环与输出","不理解赛道语义")])
body("视觉工程现有6组宿主机测试，覆盖三轮混控、伪红外、循迹、避障、球体视觉和推球状态机。它们用于核对边界条件和状态跳转，属于软件实现验证，故在本章记录，不另设独立“系统测试与结果分析”章节。")

# 8 调试问题
heading("8 课题开发与调试问题分析")
body("下表按“现象—原因—修改—结果”整理主要问题。结果只描述实际观察到的改进，不虚构成功率。")
issues=[
 ("中心采样死区","偏离后突然大幅转向","内外通道同时参与误差","仅CH2/CH3给±1，外侧只触发弯道","直行微调与弯道分离"),
 ("大弯提前转向","外侧先看见黑线便强转","外侧直接映射大误差","30 ms预触发、450 ms确认","减少缓弯误判"),
 ("丢线无方向","四框全白，PD无效","图像完全失线","右120°、左240°编码器搜索","可主动重获线路"),
 ("JPEG延迟","控制与网页滞后","回调解码且帧堆积","core 1任务、双尺度、丢帧、帧龄保护","减少陈旧帧影响"),
 ("横移堵转","车轮卡住仍输出","固定动作无反馈","500 ms堵转、6000 ms超时、进度上限","异常自动停车"),
 ("固定时间避障","电量变化导致位移不一","时间不代表距离","编码器定距、连续清障确认","重复性提高"),
 ("近场视觉盲区","球接近后出画","视场下缘受限","远场对准、FAR_SPRINT、横移、直推","降低近场图像依赖"),
 ("蓝球误识别","暗轨道被当蓝球","颜色条件过宽","加入亮度、饱和度、位置、形状、梯度","候选更稳定"),
 ("STM32烧录失败","ST-LINK持续失败","模块硬件异常","万用表检查并更换模块","成功烧录"),
 ("SysTick卡死","HAL_Delay停滞","工程缺少常规中断文件","显式SysTick_Handler调用HAL_IncTick","节拍恢复"),
 ("IMU地址确认","初始化失败难定位","接线/地址不明确","I²C扫描和WHO_AM_I","确认0x68"),
 ("串口乱码","ZS-040字符不可读","波特率不一致","两端统一9600 baud","字符正常"),
 ("BLE主从问题","两个模块无法传统主从","模块实际支持BLE模式","单ZS-040+ESP32 GATT客户端","无线链路成功"),
 ("手势抖动","阈值附近反复切换","噪声、手抖、单阈值","α=0.3、30°/20°、优势轴、3帧","动作稳定")]
table(["问题","现象","原因","修改","结果"],issues)

# 9 创新点
heading("9 系统创新点")
innov=[
 ("9.1 摄像头伪红外","四个采样框把二维图像转换为与物理红外一致的四位掩码，复用PD、弯道和回线控制器。创新点不是视觉循迹本身，而是统一感知接口及可解释的采样/阈值信息。"),
 ("9.2 正常转弯与丢线转向双机制","外侧预触发与时间窗确认负责仍能看见线路的弯道；编码器限定角度的双向搜索负责四路全白后的重获，两者均含稳定确认与安全超时。"),
 ("9.3 视觉流水线实时性优化","双尺度解码兼顾速度与精度，USB回调和JPEG解码解耦，core 1专门处理，忙时丢帧与帧龄保护以低延迟优先。最终版本不以50 fps作固定指标。"),
 ("9.4 避障防锁死机制","连续触发/清障确认、500 ms测距时效、编码器进度、阶段超时、堵转超时与最大进度共同覆盖传感器过期、车轮卡死和状态无法结束。"),
 ("9.5 推球远场冲刺射门","FAR_SPRINT在远场完成球—洞对准，再按纵向位置生成编码器冲刺距离；近场只横移，最后直推，组合视觉和里程信息解决近场盲区。"),
 ("9.6 Wi-Fi画面及详细信息显示","手机端同时显示MJPEG、四框采样、阈值、掩码、收帧/解码/丢帧、帧龄、球/洞口位置和推球状态，使“看见什么”与“为何运动”能够对应。")]
for h,t in innov: heading(h,2); body(t)
heading("9.7 STM32无线体感控制整体创新",2)
body("STM32、MPU6050、ZS-040 BLE与ESP32-S3组成自主开发的人机交互系统。低通—滞回—优势轴—连续帧防抖—周期心跳—失联停车—LCD反馈形成完整可靠性链，与老师统一提供的常规底盘硬件明确区分。")

# 10 日志
heading("10 工作日志")
logs=[
 ("8月24日下午","配置代码环境；成功烧录测试文件。","全体"),
 ("8月25日上午","连接并测试电机驱动，车轮正常转动。","全体；何晓瑜、王铜锌侧重硬件"),
 ("8月25日下午","实现红外黑白识别；使用串口助手；铺设赛道尝试循迹，掌握监测调试方法。","王润宇主导，全体参与"),
 ("8月26日上午","在GitHub搭建协作平台；实现基本循迹，转向待优化。","王润宇主导，全体参与"),
 ("8月26日下午","开发第二版循迹；第一版微调仍有问题。","王润宇"),
 ("8月27日上午","优化两版代码，对比后第二版表现更好。","王润宇"),
 ("8月27日下午","开发两版超声波代码；第一版接收异常，第二版有反应但需优化。","王润宇主导，硬件协助"),
 ("8月28日上午","基本结合循迹与避障，横移和直行待优化。","王润宇主导，全体联调"),
 ("8月28日下午","实现循迹+避障；尝试显示屏。","全体"),
 ("8月31日上午","UART烧录循迹避障；显示电机转速与距离，验收成功；摄像头初测失败。","全体"),
 ("8月31日下午","Wi-Fi手机显示摄像头；拍摄正常，黑线识别差。","王润宇主导"),
 ("9月1日上午","手机实时显示画面、黑线识别和电机指令。","王润宇"),
 ("9月1日下午","尝试六横线扫线，误差大，多版仍失败。","王润宇"),
 ("9月2日上午","可循迹直行，转弯与微调不理想。","王润宇"),
 ("9月2日下午","放弃扫线，转向同水平四框伪红外。","王润宇"),
 ("9月3日上午","伪红外初版易丢线，传输帧率低且滞后。","王润宇"),
 ("9月3日下午","优化后当时开发状态观察到50帧以上；直行和直角弯良好，大弯仍会提前转。","王润宇主导，全体调试"),
 ("9月4日上午","开始识球推球；直角轨道会被误认大球，识球不准。","何晓瑜"),
 ("9月4日下午","摄像头循迹成功；红球可识别，蓝球不可识别。","王润宇、何晓瑜"),
 ("9月7日上午","识别红球和洞口并推球，灵敏度精度不足；启动创新任务。","何晓瑜主导，全体"),
 ("9月7日下午","红球推入正常，蓝球较差；STM环境完成并开始ST-LINK测试。","何晓瑜；王铜锌"),
 ("9月8日上午","蓝球可识别但不稳；ST-LINK失败，万用表发现模块故障，更换后烧录成功。","何晓瑜；王铜锌主导硬件"),
 ("9月8日下午","蓝球识别稳定，推球继续优化；USB-TTL、STM32、ST-LINK、LQ6050通过，电脑可见方位变化。","何晓瑜；王铜锌"),
 ("9月9日上午","实现识别并推蓝球，继续提高成功率；尝试两个ZS-040主从。","何晓瑜；王铜锌"),
 ("9月9日下午","任务二验收；ZS-040无法设置传统主从。","全体；王铜锌"),
 ("9月10日上午","撰写PPT；确认BLE模式，改用一个ZS-040实现传输。","全体；王铜锌"),
 ("9月10日下午","撰写PPT；实现LQ6050+ZS-040体感控制，屏幕显示三路电机与小车状态。","全体")]
table(["日期/时段","工作内容","主要参与"],logs,[2.5,10,4])

# 11 总结与附录
heading("11 总结与展望")
body("本课题完成三层演进：以物理红外和超声波建立循迹避障与编码器闭环基线；以伪红外接口复用控制器并加入找球、洞口和十四状态推球；以STM32和MPU6050构建手持端，经BLE实现体感控制。各阶段都加入数据时效、连续确认、堵转超时、通信看门狗和状态显示，使异常能够解释并安全停止。")
body("后续可加入颜色标定、透视校正或轻量模型改善强反光和阴影；增加入洞二次确认，弥补末段定时/里程控制不足；对IMU进行陀螺零偏标定并采用互补或卡尔曼融合，改善快速运动；将体感横滚映射到后轮参与的全向横移，并统一接入自主工程的三轮闭环。还应在多场地记录成功率、平均耗时和光照条件，形成定量评价。")
doc.add_page_break(); heading("附录A 完整源代码说明")
body("完整源码以独立ZIP附件提交，包含下列四套自研工程。附件应排除build、.pio/build、managed_components、目标文件、固件和其他可再生内容。正文只列代表性代码，宏定义、状态枚举、驱动接口和测试代码以附件为准。")
table(["目录","内容","重点文件"],[
 ("lnf_avoid","物理红外循迹与超声波避障","循迹、避障、运动、编码器与显示源码"),
 ("vison_lnf_avoid","视觉循迹、识球与推球","camera_vision、pseudo_infrared、ball_vision、ball_push"),
 ("STM32_LQ6050","STM32体感手持端","platformio.ini、src/main.c"),
 ("ESP32_Car_Arduino","车端BLE接收和LCD","platformio.ini、src/main.cpp")])
body("源码附件文件名：待填写。")
heading("附录B 作品视频")
table(["项目","内容"],[("视频文件或链接","待填写"),("实车照片","待填写"),("建议内容","循迹、避障、红蓝球推入、Wi-Fi状态页、体感控制和失联停车")])

for p in doc.paragraphs: p.paragraph_format.widow_control=True
for t in doc.tables:
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
doc.core_properties.title="基于ESP32-S3的视觉自主导航与无线体感控制智能小车设计"
doc.core_properties.author="王润宇、何晓瑜、王铜锌"
doc.core_properties.subject="电子技术课程设计总结报告"
doc.save(OUT)
print(OUT)
